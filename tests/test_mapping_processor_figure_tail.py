from __future__ import annotations

import json
import zipfile
from pathlib import Path

from docx import Document as DocxDocument
from openpyxl import Workbook

from modules.mapping_processor import process_mapping_excel


HEADERS = [
    "檔案名稱/資料夾名稱/文字內容",
    "擷取段落/操作",
    "類型",
    "包含標題",
    "檔案路徑",
    "檔案名稱",
    "模板文件",
    "插入段落名稱/目的資料夾名稱",
]


def _write_mapping(path: Path, rows: list[list[str]]) -> None:
    wb = Workbook()
    ws = wb.active
    ws.append(HEADERS)
    for row in rows:
        ws.append(row)
    wb.save(path)


def _run_validate_mapping(
    tmp_path: Path,
    operation: str,
    item_type: str = "",
    include_title: str = "",
    source_name: str = "source.docx",
    source_files: list[str] | None = None,
    *,
    validate_extract_only: bool = False,
) -> tuple[dict, dict]:
    files_dir = tmp_path / "files"
    out_dir = tmp_path / "output"
    log_dir = tmp_path / "logs"
    files_dir.mkdir(parents=True, exist_ok=True)
    (out_dir / "out").mkdir(parents=True, exist_ok=True)
    log_dir.mkdir(parents=True, exist_ok=True)

    for rel_path in source_files or ["source.docx"]:
        src = files_dir / rel_path
        src.parent.mkdir(parents=True, exist_ok=True)
        src.write_text("dummy", encoding="utf-8")

    mapping_path = tmp_path / "mapping.xlsx"
    rows = [[source_name, operation, item_type, include_title, "out", "result.docx", "", ""]]
    _write_mapping(mapping_path, rows)

    result = process_mapping_excel(
        str(mapping_path),
        str(files_dir),
        str(out_dir),
        log_dir=str(log_dir),
        validate_only=not validate_extract_only,
        validate_extract_only=validate_extract_only,
    )

    log_data = {"messages": [], "runs": []}
    log_file = result.get("log_file")
    if log_file:
        with open(log_dir / log_file, "r", encoding="utf-8") as f:
            log_data = json.load(f)
    return result, log_data


def _first_step_params(log_data: dict) -> dict:
    runs = log_data.get("runs") or []
    assert runs, "expected at least one run entry"
    workflow_log = runs[0].get("workflow_log") or []
    assert workflow_log, "expected at least one workflow log entry"
    return (workflow_log[0].get("params") or {})


def _first_step_type(log_data: dict) -> str:
    runs = log_data.get("runs") or []
    assert runs, "expected at least one run entry"
    workflow_log = runs[0].get("workflow_log") or []
    assert workflow_log, "expected at least one workflow log entry"
    return str(workflow_log[0].get("type") or "")


def test_mapping_blank_type_defaults_to_extract_chapter(tmp_path: Path) -> None:
    result, log_data = _run_validate_mapping(tmp_path, "1.1 General description", item_type="")
    assert result.get("log_file") == "mapping_log.json"
    assert _first_step_type(log_data) == "extract_word_chapter"
    runs = log_data.get("runs") or []
    assert runs[0].get("workflow_log")
    assert "steps" not in runs[0]
    assert "validation_log" not in runs[0]
    params = _first_step_params(log_data)
    assert params.get("target_chapter_section") == "1.1"
    assert not any("ERROR:" in msg for msg in result.get("logs", []))


def test_mapping_chapter_range_sets_explicit_end_number(tmp_path: Path) -> None:
    result, log_data = _run_validate_mapping(tmp_path, "1.1.1-1.1.3 General description")
    assert _first_step_type(log_data) == "extract_word_chapter"
    params = _first_step_params(log_data)
    assert params.get("target_chapter_section") == "1.1.1"
    assert params.get("explicit_end_number") == "1.1.3"
    assert params.get("target_chapter_title") == "General description"
    assert not any("ERROR:" in msg for msg in result.get("logs", []))


def test_mapping_chapter_range_with_subheading_sets_subtitle(tmp_path: Path) -> None:
    result, log_data = _run_validate_mapping(
        tmp_path,
        r"1.1.1-1.1.3 General description\Device trade name",
    )
    assert _first_step_type(log_data) == "extract_word_chapter"
    params = _first_step_params(log_data)
    assert params.get("target_chapter_section") == "1.1.1"
    assert params.get("explicit_end_number") == "1.1.3"
    assert params.get("target_chapter_title") == "General description"
    assert params.get("target_subtitle") == "Device trade name"
    assert params.get("use_chapter_title") is True
    assert not any("ERROR:" in msg for msg in result.get("logs", []))


def test_mapping_chapter_range_with_start_and_end_titles_sets_both_titles(tmp_path: Path) -> None:
    result, log_data = _run_validate_mapping(
        tmp_path,
        "1.1.1 General description - 1.2.2 Intended users",
    )
    assert _first_step_type(log_data) == "extract_word_chapter"
    params = _first_step_params(log_data)
    assert params.get("target_chapter_section") == "1.1.1"
    assert params.get("target_chapter_title") == "General description"
    assert params.get("explicit_end_number") == "1.2.2"
    assert params.get("explicit_end_title") == "Intended users"
    assert not any("ERROR:" in msg for msg in result.get("logs", []))


def test_mapping_type_add_text_blank_operation_creates_insert_text(tmp_path: Path) -> None:
    result, log_data = _run_validate_mapping(
        tmp_path,
        "",
        item_type="Add Text",
        source_name="這是一段說明文字",
        source_files=[],
    )
    assert _first_step_type(log_data) == "insert_text"
    params = _first_step_params(log_data)
    assert params.get("text") == "這是一段說明文字"
    assert not any("ERROR:" in msg for msg in result.get("logs", []))


def test_mapping_type_add_text_rejects_other_operation(tmp_path: Path) -> None:
    result, log_data = _run_validate_mapping(
        tmp_path,
        "1.1 General description",
        item_type="Add Text",
        source_name="這是一段說明文字",
        source_files=[],
    )
    assert any("類型 Add Text 時，操作欄僅支援留白或 Add Text" in msg for msg in result.get("logs", []))
    runs = log_data.get("runs") or []
    assert runs
    assert all(not (run.get("workflow_log") or []) for run in runs)


def test_mapping_pdf_image_blank_operation_creates_pdf_step(tmp_path: Path) -> None:
    result, log_data = _run_validate_mapping(
        tmp_path,
        "",
        item_type="PDF Image",
        source_name="source.pdf",
        source_files=["source.pdf"],
    )
    assert _first_step_type(log_data) == "extract_pdf_pages_as_images"
    params = _first_step_params(log_data)
    assert Path(str(params.get("input_file", ""))).name == "source.pdf"
    assert not any("ERROR:" in msg for msg in result.get("logs", []))


def test_mapping_pdf_image_requires_pdf_file(tmp_path: Path) -> None:
    result, log_data = _run_validate_mapping(
        tmp_path,
        "All Pages",
        item_type="PDF Image",
        source_name="source.docx",
        source_files=["source.docx"],
    )
    assert any("PDF Image 類型僅支援 PDF 檔案" in msg for msg in result.get("logs", []))
    runs = log_data.get("runs") or []
    assert runs
    assert all(not (run.get("workflow_log") or []) for run in runs)


def test_mapping_copy_file_blank_operation_creates_copy_file_step(tmp_path: Path) -> None:
    result, log_data = _run_validate_mapping(
        tmp_path,
        "",
        item_type="Copy File",
        source_name="source.pdf",
        source_files=["source.pdf"],
    )
    assert _first_step_type(log_data) == "copy_file"
    params = _first_step_params(log_data)
    assert Path(str(params.get("source", ""))).name == "source.pdf"
    assert params.get("keywords") == ""
    assert not any("ERROR:" in msg for msg in result.get("logs", []))


def test_mapping_copy_file_keywords_create_copy_file_step(tmp_path: Path) -> None:
    files_dir = tmp_path / "files"
    out_dir = tmp_path / "output"
    log_dir = tmp_path / "logs"
    files_dir.mkdir(parents=True, exist_ok=True)
    out_dir.mkdir(parents=True, exist_ok=True)
    log_dir.mkdir(parents=True, exist_ok=True)
    (files_dir / "RootA").mkdir(parents=True, exist_ok=True)

    mapping_path = tmp_path / "mapping.xlsx"
    rows = [["RootA", "IFU,EO", "Copy File", "", "out", "", "", ""]]
    _write_mapping(mapping_path, rows)

    result = process_mapping_excel(
        str(mapping_path),
        str(files_dir),
        str(out_dir),
        log_dir=str(log_dir),
        validate_only=True,
    )

    log_file = result.get("log_file")
    assert log_file == "mapping_log.json"
    with open(log_dir / log_file, "r", encoding="utf-8") as f:
        log_data = json.load(f)
    assert _first_step_type(log_data) == "copy_file"
    params = _first_step_params(log_data)
    assert Path(str(params.get("source", ""))).name == "RootA"
    assert params.get("keywords") == "IFU,EO"
    assert not any("ERROR:" in msg for msg in result.get("logs", []))


def test_mapping_copy_folder_blank_operation_creates_copy_folder_step(tmp_path: Path) -> None:
    files_dir = tmp_path / "files"
    out_dir = tmp_path / "output"
    log_dir = tmp_path / "logs"
    files_dir.mkdir(parents=True, exist_ok=True)
    out_dir.mkdir(parents=True, exist_ok=True)
    log_dir.mkdir(parents=True, exist_ok=True)
    (files_dir / "FolderA" / "IFU").mkdir(parents=True, exist_ok=True)

    mapping_path = tmp_path / "mapping.xlsx"
    rows = [["FolderA/IFU", "", "Copy Folder", "", "out", "", "", ""]]
    _write_mapping(mapping_path, rows)

    result = process_mapping_excel(
        str(mapping_path),
        str(files_dir),
        str(out_dir),
        log_dir=str(log_dir),
        validate_only=True,
    )

    log_file = result.get("log_file")
    assert log_file == "mapping_log.json"
    assert log_file
    with open(log_dir / log_file, "r", encoding="utf-8") as f:
        log_data = json.load(f)
    assert _first_step_type(log_data) == "copy_folder"
    params = _first_step_params(log_data)
    assert Path(str(params.get("source", ""))).name == "IFU"
    assert params.get("keywords") == ""
    assert not any("ERROR:" in msg for msg in result.get("logs", []))


def test_mapping_copy_folder_keywords_create_copy_folder_step(tmp_path: Path) -> None:
    files_dir = tmp_path / "files"
    out_dir = tmp_path / "output"
    log_dir = tmp_path / "logs"
    files_dir.mkdir(parents=True, exist_ok=True)
    out_dir.mkdir(parents=True, exist_ok=True)
    log_dir.mkdir(parents=True, exist_ok=True)
    (files_dir / "RootA").mkdir(parents=True, exist_ok=True)

    mapping_path = tmp_path / "mapping.xlsx"
    rows = [["RootA", "IFU,Label", "Copy Folder", "", "out", "", "", ""]]
    _write_mapping(mapping_path, rows)

    result = process_mapping_excel(
        str(mapping_path),
        str(files_dir),
        str(out_dir),
        log_dir=str(log_dir),
        validate_only=True,
    )

    log_file = result.get("log_file")
    assert log_file == "mapping_log.json"
    with open(log_dir / log_file, "r", encoding="utf-8") as f:
        log_data = json.load(f)
    assert _first_step_type(log_data) == "copy_folder"
    params = _first_step_params(log_data)
    assert Path(str(params.get("source", ""))).name == "RootA"
    assert params.get("keywords") == "IFU,Label"
    assert not any("ERROR:" in msg for msg in result.get("logs", []))


def test_mapping_figure_tail_title(tmp_path: Path) -> None:
    result, log_data = _run_validate_mapping(tmp_path, "1.1 Figure 1|title=Overview Figure")
    params = _first_step_params(log_data)
    assert params.get("target_caption_label") == "Figure 1"
    assert params.get("target_figure_title") == "Overview Figure"
    assert "target_figure_index" not in params
    assert not any("ERROR:" in msg for msg in result.get("logs", []))


def test_mapping_figure_tail_index(tmp_path: Path) -> None:
    result, log_data = _run_validate_mapping(tmp_path, "1.1 Figure 1|index=2")
    params = _first_step_params(log_data)
    assert params.get("target_caption_label") == "Figure 1"
    assert params.get("target_figure_index") == 2
    assert not any("ERROR:" in msg for msg in result.get("logs", []))


def test_mapping_figure_tail_title_and_index(tmp_path: Path) -> None:
    result, log_data = _run_validate_mapping(tmp_path, "1.1 Figure 1|title=System Overview|index=3")
    params = _first_step_params(log_data)
    assert params.get("target_figure_title") == "System Overview"
    assert params.get("target_figure_index") == 3
    assert not any("ERROR:" in msg for msg in result.get("logs", []))


def test_mapping_tail_requires_figure_or_table_label(tmp_path: Path) -> None:
    result, log_data = _run_validate_mapping(tmp_path, "1.1 Chapter|title=No Label")
    assert any("使用 title/index 參數時必須指定 Figure 或 Table 標籤" in msg for msg in result.get("logs", []))
    runs = log_data.get("runs") or []
    assert runs
    assert all(not (run.get("workflow_log") or []) for run in runs)


def test_mapping_type_figure_allows_tail_without_label(tmp_path: Path) -> None:
    result, log_data = _run_validate_mapping(
        tmp_path,
        "1.1 General description|title=System overview|index=2",
        item_type="Figure",
    )
    assert _first_step_type(log_data) == "extract_specific_figure_from_word"
    params = _first_step_params(log_data)
    assert params.get("target_caption_label") == ""
    assert params.get("target_figure_title") == "System overview"
    assert params.get("target_figure_index") == 2
    assert not any("ERROR:" in msg for msg in result.get("logs", []))


def test_mapping_type_table_allows_tail_without_label(tmp_path: Path) -> None:
    result, log_data = _run_validate_mapping(
        tmp_path,
        "1.1 General description|title=Table summary|index=1",
        item_type="Table",
    )
    assert _first_step_type(log_data) == "extract_specific_table_from_word"
    params = _first_step_params(log_data)
    assert params.get("target_caption_label") == ""
    assert params.get("target_table_title") == "Table summary"
    assert params.get("target_table_index") == 1
    assert not any("ERROR:" in msg for msg in result.get("logs", []))


def test_mapping_type_conflicts_with_label(tmp_path: Path) -> None:
    result, log_data = _run_validate_mapping(tmp_path, "1.1 Table 1", item_type="Figure")
    assert any("類型欄位與操作內容不一致" in msg for msg in result.get("logs", []))
    runs = log_data.get("runs") or []
    assert runs
    assert all(not (run.get("workflow_log") or []) for run in runs)


def test_mapping_type_figure_requires_caption_or_title_or_index(tmp_path: Path) -> None:
    result, log_data = _run_validate_mapping(tmp_path, "1.1 General description", item_type="Figure")
    assert any("Figure 擷取至少需提供 caption、title 或 index 其中之一" in msg for msg in result.get("logs", []))
    runs = log_data.get("runs") or []
    assert runs
    assert all(not (run.get("workflow_log") or []) for run in runs)


def test_mapping_figure_index_validation(tmp_path: Path) -> None:
    bad_result, _ = _run_validate_mapping(tmp_path / "bad", "1.1 Figure 1|index=abc")
    zero_result, _ = _run_validate_mapping(tmp_path / "zero", "1.1 Figure 1|index=0")
    assert any("index 必須是正整數: abc" in msg for msg in bad_result.get("logs", []))
    assert any("index 必須大於 0: 0" in msg for msg in zero_result.get("logs", []))


def test_mapping_table_tail_behavior_unchanged(tmp_path: Path) -> None:
    result, log_data = _run_validate_mapping(tmp_path, "1.1 Table 1|title=Table Name|index=4")
    params = _first_step_params(log_data)
    assert params.get("target_caption_label") == "Table 1"
    assert params.get("target_table_title") == "Table Name"
    assert params.get("target_table_index") == 4
    assert "target_figure_title" not in params
    assert "target_figure_index" not in params
    assert not any("ERROR:" in msg for msg in result.get("logs", []))


def test_mapping_chapter_include_title_false_sets_hide_flag(tmp_path: Path) -> None:
    result, log_data = _run_validate_mapping(
        tmp_path,
        "1.1 General description",
        include_title="FALSE",
    )
    params = _first_step_params(log_data)
    assert params.get("hide_chapter_title") is True
    assert not any("ERROR:" in msg for msg in result.get("logs", []))


def test_mapping_figure_include_title_false_disables_caption(tmp_path: Path) -> None:
    result, log_data = _run_validate_mapping(
        tmp_path,
        "1.1 Figure 1|title=Overview Figure",
        include_title="否",
    )
    params = _first_step_params(log_data)
    assert params.get("include_caption") is False
    assert not any("ERROR:" in msg for msg in result.get("logs", []))


def test_mapping_table_include_title_false_disables_caption(tmp_path: Path) -> None:
    result, log_data = _run_validate_mapping(
        tmp_path,
        "1.1 Table 1|title=Table Name",
        include_title="N",
    )
    params = _first_step_params(log_data)
    assert params.get("include_caption") is False
    assert not any("ERROR:" in msg for msg in result.get("logs", []))


def test_mapping_include_title_invalid_value_errors(tmp_path: Path) -> None:
    result, log_data = _run_validate_mapping(
        tmp_path,
        "1.1 General description",
        include_title="maybe",
    )
    assert any("包含標題欄位值無效: maybe" in msg for msg in result.get("logs", []))
    runs = log_data.get("runs") or []
    assert not runs or all(not (run.get("workflow_log") or []) for run in runs)


def test_mapping_source_relative_path_resolves_file(tmp_path: Path) -> None:
    result, log_data = _run_validate_mapping(
        tmp_path,
        "1.1 General description",
        source_name="FolderA/source.docx",
        source_files=["FolderA/source.docx"],
    )
    params = _first_step_params(log_data)
    assert Path(str(params.get("input_file", ""))).name == "source.docx"
    assert "FolderA" in str(params.get("input_file", ""))
    assert not any("ERROR:" in msg for msg in result.get("logs", []))


def test_mapping_duplicate_filename_requires_relative_path(tmp_path: Path) -> None:
    result, log_data = _run_validate_mapping(
        tmp_path,
        "1.1 General description",
        source_name="source.docx",
        source_files=["FolderA/source.docx", "FolderB/source.docx"],
    )
    assert any("找到多個名稱相同檔案 source.docx" in msg for msg in result.get("logs", []))
    runs = log_data.get("runs") or []
    assert runs
    assert all(not (run.get("workflow_log") or []) for run in runs)


def test_mapping_blank_output_path_defaults_to_output_root(tmp_path: Path) -> None:
    files_dir = tmp_path / "files"
    out_dir = tmp_path / "output"
    log_dir = tmp_path / "logs"
    files_dir.mkdir(parents=True, exist_ok=True)
    out_dir.mkdir(parents=True, exist_ok=True)
    log_dir.mkdir(parents=True, exist_ok=True)

    src = files_dir / "source.docx"
    src.write_text("dummy", encoding="utf-8")

    mapping_path = tmp_path / "mapping.xlsx"
    rows = [["source.docx", "1.1 General description", "", "", "", "result.docx", "", ""]]
    _write_mapping(mapping_path, rows)

    result = process_mapping_excel(
        str(mapping_path),
        str(files_dir),
        str(out_dir),
        log_dir=str(log_dir),
        validate_only=True,
    )

    assert not any("缺少輸出路徑" in msg for msg in result.get("logs", []))
    log_file = result.get("log_file")
    assert log_file
    with open(log_dir / log_file, "r", encoding="utf-8") as f:
        log_data = json.load(f)
    runs = log_data.get("runs") or []
    assert runs
    assert runs[0].get("output") == "result.docx"


def test_mapping_outputs_are_packaged_into_zip(tmp_path: Path, monkeypatch) -> None:
    files_dir = tmp_path / "files"
    out_dir = tmp_path / "output"
    log_dir = tmp_path / "logs"
    files_dir.mkdir(parents=True, exist_ok=True)
    out_dir.mkdir(parents=True, exist_ok=True)
    log_dir.mkdir(parents=True, exist_ok=True)

    src = files_dir / "source.docx"
    src.write_text("dummy", encoding="utf-8")

    mapping_path = tmp_path / "mapping.xlsx"
    rows = [
        ["source.docx", "1.1 General description", "", "", "pkg/A", "a.docx", "", ""],
        ["source.docx", "1.2 General description", "", "", "pkg/B", "b.docx", "", ""],
    ]
    _write_mapping(mapping_path, rows)

    def fake_run_workflow(steps, workdir, template=None):
        result_docx = Path(workdir) / "result.docx"
        doc = DocxDocument()
        doc.add_paragraph(f"generated-{len(steps)}")
        doc.save(result_docx)
        return {"result_docx": str(result_docx), "log_json": []}

    monkeypatch.setattr("modules.mapping_processor.run_workflow", fake_run_workflow)
    monkeypatch.setattr("modules.mapping_processor.apply_basic_style", lambda *args, **kwargs: True)
    monkeypatch.setattr("modules.mapping_processor.remove_hidden_runs", lambda *args, **kwargs: True)
    monkeypatch.setattr("modules.mapping_processor.hide_paragraphs_with_text", lambda *args, **kwargs: True)

    result = process_mapping_excel(
        str(mapping_path),
        str(files_dir),
        str(out_dir),
        log_dir=str(log_dir),
        validate_only=False,
    )

    zip_file = result.get("zip_file")
    assert zip_file
    zip_path = out_dir / zip_file
    assert zip_path.is_file()
    with zipfile.ZipFile(zip_path, "r") as zf:
        names = sorted(zf.namelist())
    assert names == ["pkg/A/a.docx", "pkg/B/b.docx"]


def test_mapping_copy_outputs_use_conflict_suffix_and_zip(tmp_path: Path) -> None:
    files_dir = tmp_path / "files"
    out_dir = tmp_path / "output"
    log_dir = tmp_path / "logs"
    files_dir.mkdir(parents=True, exist_ok=True)
    out_dir.mkdir(parents=True, exist_ok=True)
    log_dir.mkdir(parents=True, exist_ok=True)

    knee_ifu = (
        files_dir
        / "輸入-測試路徑"
        / "TD-III-011-USTAR II Knee System"
        / "Section 2_Information Supplied by the Manufacturer"
        / "IFU"
    )
    hip_ifu = (
        files_dir
        / "輸入-測試路徑"
        / "TD-III-012-USTAR II Hip System"
        / "Section 2_Information Supplied by the Manufacturer"
        / "IFU"
    )
    knee_ifu.mkdir(parents=True, exist_ok=True)
    hip_ifu.mkdir(parents=True, exist_ok=True)
    (knee_ifu / "knee.txt").write_text("knee", encoding="utf-8")
    (hip_ifu / "hip.txt").write_text("hip", encoding="utf-8")
    (files_dir / "輸入-測試路徑" / "TD-III-011-USTAR II Knee System" / "labeling.pdf").write_text("knee pdf", encoding="utf-8")
    (files_dir / "輸入-測試路徑" / "TD-III-012-USTAR II Hip System" / "labeling.pdf").write_text("hip pdf", encoding="utf-8")

    mapping_path = tmp_path / "mapping.xlsx"
    rows = [
        [
            "輸入-測試路徑/TD-III-011-USTAR II Knee System/Section 2_Information Supplied by the Manufacturer/IFU",
            "",
            "Copy Folder",
            "",
            "pkg/folders",
            "",
            "",
            "",
        ],
        [
            "輸入-測試路徑/TD-III-012-USTAR II Hip System/Section 2_Information Supplied by the Manufacturer/IFU",
            "",
            "Copy Folder",
            "",
            "pkg/folders",
            "",
            "",
            "",
        ],
        [
            "輸入-測試路徑/TD-III-011-USTAR II Knee System/labeling.pdf",
            "",
            "Copy File",
            "",
            "pkg/files",
            "",
            "",
            "",
        ],
        [
            "輸入-測試路徑/TD-III-012-USTAR II Hip System/labeling.pdf",
            "",
            "Copy File",
            "",
            "pkg/files",
            "",
            "",
            "",
        ],
    ]
    _write_mapping(mapping_path, rows)

    result = process_mapping_excel(
        str(mapping_path),
        str(files_dir),
        str(out_dir),
        log_dir=str(log_dir),
        validate_only=False,
    )

    assert result.get("log_file") == "mapping_log.json"
    assert sorted(Path(p).name for p in result.get("outputs", [])) == ["labeling_hip.pdf", "labeling_knee.pdf"]
    zip_file = result.get("zip_file")
    assert zip_file
    with zipfile.ZipFile(out_dir / zip_file, "r") as zf:
        names = sorted(zf.namelist())
    assert names == [
        "pkg/files/labeling_hip.pdf",
        "pkg/files/labeling_knee.pdf",
        "pkg/folders/IFU_hip/hip.txt",
        "pkg/folders/IFU_knee/knee.txt",
    ]


def test_mapping_copy_folder_keywords_copy_matching_subfolders(tmp_path: Path) -> None:
    files_dir = tmp_path / "files"
    out_dir = tmp_path / "output"
    log_dir = tmp_path / "logs"
    files_dir.mkdir(parents=True, exist_ok=True)
    out_dir.mkdir(parents=True, exist_ok=True)
    log_dir.mkdir(parents=True, exist_ok=True)

    knee_ifu = files_dir / "source_root" / "Knee System" / "IFU"
    hip_ifu = files_dir / "source_root" / "Hip System" / "IFU"
    other_folder = files_dir / "source_root" / "Misc System" / "Label sample"
    knee_ifu.mkdir(parents=True, exist_ok=True)
    hip_ifu.mkdir(parents=True, exist_ok=True)
    other_folder.mkdir(parents=True, exist_ok=True)
    (knee_ifu / "knee.txt").write_text("knee", encoding="utf-8")
    (hip_ifu / "hip.txt").write_text("hip", encoding="utf-8")
    (other_folder / "label.txt").write_text("label", encoding="utf-8")

    mapping_path = tmp_path / "mapping.xlsx"
    rows = [["source_root", "IFU", "Copy Folder", "", "pkg/folders", "", "", ""]]
    _write_mapping(mapping_path, rows)

    result = process_mapping_excel(
        str(mapping_path),
        str(files_dir),
        str(out_dir),
        log_dir=str(log_dir),
        validate_only=False,
    )

    assert result.get("log_file") == "mapping_log.json"
    zip_file = result.get("zip_file")
    assert zip_file
    with zipfile.ZipFile(out_dir / zip_file, "r") as zf:
        names = sorted(zf.namelist())
    assert names == [
        "pkg/folders/IFU_hip/hip.txt",
        "pkg/folders/IFU_knee/knee.txt",
    ]


def test_mapping_copy_file_keywords_copy_matching_files(tmp_path: Path) -> None:
    files_dir = tmp_path / "files"
    out_dir = tmp_path / "output"
    log_dir = tmp_path / "logs"
    files_dir.mkdir(parents=True, exist_ok=True)
    out_dir.mkdir(parents=True, exist_ok=True)
    log_dir.mkdir(parents=True, exist_ok=True)

    source_root = files_dir / "source_root"
    source_root.mkdir(parents=True, exist_ok=True)
    (source_root / "Shipping simulation test EO.pdf").write_text("eo", encoding="utf-8")
    (source_root / "Shipping simulation test Gamma.pdf").write_text("gamma", encoding="utf-8")
    (source_root / "Other file.pdf").write_text("other", encoding="utf-8")

    mapping_path = tmp_path / "mapping.xlsx"
    rows = [["source_root", "Shipping simulation test,EO", "Copy File", "", "pkg/files", "", "", ""]]
    _write_mapping(mapping_path, rows)

    result = process_mapping_excel(
        str(mapping_path),
        str(files_dir),
        str(out_dir),
        log_dir=str(log_dir),
        validate_only=False,
    )

    assert result.get("log_file") == "mapping_log.json"
    zip_file = result.get("zip_file")
    assert zip_file
    with zipfile.ZipFile(out_dir / zip_file, "r") as zf:
        names = sorted(zf.namelist())
    assert names == ["pkg/files/Shipping simulation test EO.pdf"]


def test_mapping_validate_extract_only_runs_workflow_validation(tmp_path: Path, monkeypatch) -> None:
    result, log_data = None, None

    def fake_run_workflow(steps, workdir, template=None):
        return {
            "result_docx": str(Path(workdir) / "result.docx"),
            "log_json": [
                {
                    "step": 1,
                    "type": "extract_word_chapter",
                    "params": steps[0]["params"],
                    "status": "error",
                    "error": "No content extracted",
                }
            ],
        }

    monkeypatch.setattr("modules.mapping_processor.run_workflow", fake_run_workflow)

    result, log_data = _run_validate_mapping(
        tmp_path,
        "1.1 General description",
        validate_extract_only=True,
    )

    runs = log_data.get("runs") or []
    assert runs
    workflow_log = runs[0].get("workflow_log") or []
    assert workflow_log
    assert workflow_log[0].get("status") == "error"
    assert workflow_log[0].get("error") == "No content extracted"
    assert result.get("outputs") == []
    assert result.get("zip_file") is None
